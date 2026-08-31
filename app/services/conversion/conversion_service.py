"""PDF Conversion service - handles all format conversions."""
import os
from pathlib import Path
from typing import Callable, Optional

import fitz  # PyMuPDF
from PIL import Image


class ConversionService:
    """Service for converting between PDF and other formats."""

    def convert_batch(
        self,
        files: list[str],
        mode: int,
        output_dir: Optional[str] = None,
        fmt: str = "png",
        dpi: int = 300,
        quality: int = 90,
        transparent: bool = False,
        overwrite: bool = False,
        progress_callback: Optional[Callable[[int], None]] = None,
        progress_detail_callback: Optional[Callable[[int, str], None]] = None,
        is_cancelled: Optional[Callable[[], bool]] = None,
    ) -> list[str]:
        """Batch convert files based on mode.

        Modes:
            0 - PDF to Images
            1 - Images to PDF
            2 - PDF to Word (DOCX)
            3 - Word to PDF
            4 - PDF to HTML
            5 - PDF to Text
        """
        output_files: list[str] = []
        total = len(files)

        for i, file_path in enumerate(files):
            if is_cancelled and is_cancelled():
                break

            percent = int((i / total) * 100)
            if progress_callback:
                progress_callback(percent)
            if progress_detail_callback:
                progress_detail_callback(percent, f"Processing: {Path(file_path).name}")

            try:
                if mode == 0:
                    results = self.pdf_to_images(
                        file_path, output_dir, fmt, dpi, quality, transparent, overwrite
                    )
                    output_files.extend(results)
                elif mode == 1:
                    result = self.images_to_pdf(
                        [file_path] if i == 0 else [],
                        output_dir,
                        quality,
                    )
                    if result:
                        output_files.append(result)
                elif mode == 2:
                    result = self.pdf_to_docx(file_path, output_dir)
                    if result:
                        output_files.append(result)
                elif mode == 3:
                    result = self.docx_to_pdf(file_path, output_dir)
                    if result:
                        output_files.append(result)
                elif mode == 4:
                    result = self.pdf_to_html(file_path, output_dir)
                    if result:
                        output_files.append(result)
                elif mode == 5:
                    result = self.pdf_to_text(file_path, output_dir)
                    if result:
                        output_files.append(result)
            except Exception as e:
                raise RuntimeError(f"Error processing {Path(file_path).name}: {e}") from e

        # For mode 1, batch all images into one PDF
        if mode == 1 and files:
            result = self.images_to_pdf(files, output_dir, quality)
            if result:
                output_files = [result]

        if progress_callback:
            progress_callback(100)

        return output_files

    def pdf_to_images(
        self,
        file_path: str,
        output_dir: Optional[str] = None,
        fmt: str = "png",
        dpi: int = 300,
        quality: int = 90,
        transparent: bool = False,
        overwrite: bool = False,
    ) -> list[str]:
        """Convert PDF pages to images."""
        path = Path(file_path)
        out_dir = Path(output_dir) if output_dir else path.parent
        out_dir.mkdir(parents=True, exist_ok=True)

        output_files: list[str] = []
        zoom = dpi / 72.0
        matrix = fitz.Matrix(zoom, zoom)

        doc = fitz.open(file_path)
        try:
            for page_num in range(len(doc)):
                page = doc[page_num]

                # Render with alpha if transparent
                if transparent and fmt == "png":
                    pix = page.get_pixmap(matrix=matrix, alpha=True)
                else:
                    pix = page.get_pixmap(matrix=matrix, alpha=False)

                out_name = f"{path.stem}_page_{page_num + 1}.{fmt}"
                out_path = out_dir / out_name

                if not overwrite and out_path.exists():
                    counter = 1
                    while out_path.exists():
                        out_name = f"{path.stem}_page_{page_num + 1}_{counter}.{fmt}"
                        out_path = out_dir / out_name
                        counter += 1

                if fmt in ("jpg", "jpeg"):
                    pix.save(str(out_path), jpg_quality=quality)
                elif fmt == "png":
                    pix.save(str(out_path))
                else:
                    # For WEBP, TIFF, BMP - use PIL
                    img_data = pix.tobytes("png")
                    img = Image.open(__import__("io").BytesIO(img_data))
                    if fmt == "webp":
                        img.save(str(out_path), "WEBP", quality=quality)
                    elif fmt == "tiff":
                        img.save(str(out_path), "TIFF")
                    elif fmt == "bmp":
                        img.save(str(out_path), "BMP")

                output_files.append(str(out_path))
        finally:
            doc.close()

        return output_files

    def images_to_pdf(
        self,
        image_files: list[str],
        output_dir: Optional[str] = None,
        quality: int = 90,
    ) -> str:
        """Convert multiple images to a single PDF."""
        if not image_files:
            return ""

        first_path = Path(image_files[0])
        out_dir = Path(output_dir) if output_dir else first_path.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        output_path = out_dir / f"{first_path.stem}_combined.pdf"

        doc = fitz.open()
        try:
            for img_path in image_files:
                img = Image.open(img_path)
                # Convert RGBA to RGB for PDF
                if img.mode == "RGBA":
                    bg = Image.new("RGB", img.size, (255, 255, 255))
                    bg.paste(img, mask=img.split()[3])
                    img = bg
                elif img.mode != "RGB":
                    img = img.convert("RGB")

                # Save as temp PNG for fitz
                import io
                img_bytes = io.BytesIO()
                img.save(img_bytes, format="PNG")
                img_bytes.seek(0)

                # Insert as page
                img_doc = fitz.open(stream=img_bytes.read(), filetype="png")
                rect = img_doc[0].rect
                pdf_page = doc.new_page(width=rect.width, height=rect.height)
                pdf_page.insert_image(rect, stream=img_doc[0].get_pixmap().tobytes("png"))
                img_doc.close()
                img.close()

            doc.save(str(output_path))
        finally:
            doc.close()

        return str(output_path)

    def pdf_to_text(
        self,
        file_path: str,
        output_dir: Optional[str] = None,
    ) -> str:
        """Extract text from PDF."""
        path = Path(file_path)
        out_dir = Path(output_dir) if output_dir else path.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        output_path = out_dir / f"{path.stem}.txt"

        doc = fitz.open(file_path)
        try:
            text_content = []
            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text")
                text_content.append(f"--- Page {page_num + 1} ---\n{text}\n")

            with open(output_path, "w", encoding="utf-8") as f:
                f.write("\n".join(text_content))
        finally:
            doc.close()

        return str(output_path)

    def pdf_to_html(
        self,
        file_path: str,
        output_dir: Optional[str] = None,
    ) -> str:
        """Convert PDF to HTML."""
        path = Path(file_path)
        out_dir = Path(output_dir) if output_dir else path.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        output_path = out_dir / f"{path.stem}.html"

        doc = fitz.open(file_path)
        try:
            html_parts = [
                "<!DOCTYPE html>",
                "<html><head>",
                f"<title>{path.stem}</title>",
                '<meta charset="utf-8">',
                "<style>body{font-family:sans-serif;max-width:900px;margin:auto;padding:20px;}"
                ".page{border:1px solid #ccc;padding:20px;margin:20px 0;page-break-after:always;}</style>",
                "</head><body>",
            ]

            for page_num in range(len(doc)):
                page = doc[page_num]
                html = page.get_text("html")
                html_parts.append(f'<div class="page">{html}</div>')

            html_parts.append("</body></html>")

            with open(output_path, "w", encoding="utf-8") as f:
                f.write("\n".join(html_parts))
        finally:
            doc.close()

        return str(output_path)

    def pdf_to_docx(
        self,
        file_path: str,
        output_dir: Optional[str] = None,
    ) -> str:
        """Convert PDF to DOCX (text extraction with formatting)."""
        from docx import Document
        from docx.shared import Pt, Inches

        path = Path(file_path)
        out_dir = Path(output_dir) if output_dir else path.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        output_path = out_dir / f"{path.stem}.docx"

        doc = fitz.open(file_path)
        try:
            word_doc = Document()

            for page_num in range(len(doc)):
                page = doc[page_num]
                blocks = page.get_text("dict")["blocks"]

                for block in blocks:
                    if block["type"] == 0:  # Text block
                        for line in block.get("lines", []):
                            paragraph_text = ""
                            for span in line.get("spans", []):
                                paragraph_text += span.get("text", "")
                            if paragraph_text.strip():
                                para = word_doc.add_paragraph()
                                run = para.add_run(paragraph_text)
                                # Apply approximate font size
                                if line["spans"]:
                                    font_size = line["spans"][0].get("size", 11)
                                    run.font.size = Pt(font_size)
                    elif block["type"] == 1:  # Image block
                        # Extract and insert image
                        try:
                            img_data = block.get("image")
                            if img_data:
                                import io
                                img_stream = io.BytesIO(img_data)
                                word_doc.add_picture(img_stream, width=Inches(5))
                        except Exception:
                            pass

                # Add page break between pages
                if page_num < len(doc) - 1:
                    word_doc.add_page_break()

            word_doc.save(str(output_path))
        finally:
            doc.close()

        return str(output_path)

    def docx_to_pdf(
        self,
        file_path: str,
        output_dir: Optional[str] = None,
    ) -> str:
        """Convert DOCX to PDF using ReportLab."""
        from docx import Document
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.units import inch
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet

        path = Path(file_path)
        out_dir = Path(output_dir) if output_dir else path.parent
        out_dir.mkdir(parents=True, exist_ok=True)
        output_path = out_dir / f"{path.stem}.pdf"

        word_doc = Document(file_path)
        styles = getSampleStyleSheet()

        pdf_doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            rightMargin=inch,
            leftMargin=inch,
            topMargin=inch,
            bottomMargin=inch,
        )

        story = []
        for para in word_doc.paragraphs:
            if para.text.strip():
                # Map Word styles to ReportLab styles
                style_name = "Normal"
                if para.style and para.style.name:
                    if "Heading 1" in para.style.name:
                        style_name = "Heading1"
                    elif "Heading 2" in para.style.name:
                        style_name = "Heading2"
                    elif "Title" in para.style.name:
                        style_name = "Title"

                style = styles.get(style_name, styles["Normal"])
                story.append(Paragraph(para.text, style))
                story.append(Spacer(1, 6))

        if story:
            pdf_doc.build(story)
        else:
            # Create empty PDF if no content
            pdf_doc.build([Spacer(1, 1)])

        return str(output_path)
